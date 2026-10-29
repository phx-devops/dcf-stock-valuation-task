"""Deterministic verifier for the nke-dcf-valuation task.

PASS = every test in this file passes (pytest exit code 0). Tests whose name contains
``_diag_`` are diagnostics: they are recorded in the report but can never fail the run.

Environment variables (all optional):
  APP_DIR                  default /app   (assumptions.json, market_data.json, filings/, output/)
  OUT_DIR                  default $APP_DIR/output
  GT_PATH                  default ground_truth.json next to this file (hidden from the solver)
  SOFFICE                  default soffice (pin the LibreOffice version in the Dockerfile)
  REPORT_PATH              default /logs/verifier/report.json (diagnostics only)
  ALLOW_ILLUSTRATIVE_GT=1  authoring/self-test only: lets an illustrative ground truth run
                           (page/source checks are skipped). Never set it in the real run.
"""
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import reference_model as rm  # noqa: E402

APP = Path(os.environ.get("APP_DIR", "/app"))
OUT = Path(os.environ.get("OUT_DIR", str(APP / "output")))
GT_PATH = Path(os.environ.get("GT_PATH", str(Path(__file__).resolve().parent / "ground_truth.json")))
REPORT = Path(os.environ.get("REPORT_PATH", "/logs/verifier/report.json"))
SOFFICE = os.environ.get("SOFFICE", "soffice")
XLSX = OUT / "nke_dcf.xlsx"
SUMMARY = OUT / "summary.json"

# ---- tolerances (single source of truth) -------------------------------------------------
TOL_INPUT_ABS = 1e-6      # extracted inputs must equal the reported figure
TOL_MODEL_REL = 1e-6      # every intermediate workbook row vs reference
TOL_E2E_REL = 0.01        # summary EV / equity / value per share vs reference (spec: +-1%)
TOL_WACC_ABS = 0.0010     # spec: +-10 bps
TOL_JSON_REL = 1e-6       # summary.json vs recalculated workbook (allows JSON rounding)
RECALC_TIMEOUT = 180
ALLOWED_FUNCS = {"SUM", "AVERAGE", "POWER", "IF", "INDEX", "MATCH", "ABS"}
REQUIRED_SHEETS = ["Assumptions", "Inputs", "WACC", "Historicals", "DCF"]
ASSUMPTION_KEYS = ["risk_free_rate", "equity_risk_premium", "beta", "pre_tax_cost_of_debt",
                   "terminal_growth_rate", "explicit_forecast_years", "market_price"]
ERROR_PREFIXES = ("#", "Err:")

DIAG = []


def diag(test_id, ok, detail=""):
    DIAG.append({"id": test_id, "ok": bool(ok), "detail": str(detail)})


@pytest.fixture(scope="session", autouse=True)
def _write_report():
    yield
    try:
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps({"diagnostics": DIAG}, indent=2))
    except OSError:
        pass


# ---- helpers ----------------------------------------------------------------------------
def isnum(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x)


def close(a, b, rel=TOL_MODEL_REL, abs_=1e-9):
    return abs(a - b) <= max(abs_, rel * max(abs(a), abs(b)))


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def find_row(ws, label):
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == label]
    return rows[0] if len(rows) == 1 else None


def read_outputs(wb):
    """Read every (sheet, label) in the contract from a recalculated (data_only) workbook."""
    out = {}
    for (sheet, label), ncols in rm.OUTPUT_SPEC.items():
        if sheet not in wb.sheetnames:
            out[(sheet, label)] = None
            continue
        ws = wb[sheet]
        row = find_row(ws, label)
        if row is None:
            out[(sheet, label)] = None
        elif ncols == 0:
            out[(sheet, label)] = ws.cell(row, 2).value
        else:
            out[(sheet, label)] = [ws.cell(row, 2 + i).value for i in range(ncols)]
    return out


def compare(actual, expected, rel=TOL_MODEL_REL, keys=None):
    bad = []
    for key in (keys if keys is not None else expected):
        exp, act = expected[key], actual.get(key)
        if isinstance(exp, list):
            if not isinstance(act, list) or len(act) != len(exp):
                bad.append((key, "missing or wrong shape", act))
                continue
            for i, (a, e) in enumerate(zip(act, exp)):
                if e is None:
                    continue
                if not isnum(a) or not close(a, e, rel):
                    bad.append((key, i, a, e))
        elif not isnum(act) or not close(act, exp, rel):
            bad.append((key, act, exp))
    return bad


def recalc_book(src):
    """Strip cached values (openpyxl round trip) and recalculate with LibreOffice headless."""
    work = Path(tempfile.mkdtemp(prefix="nke_recalc_"))
    (work / "in").mkdir()
    (work / "out").mkdir()
    stripped = work / "in" / "book.xlsx"
    wb = load_workbook(src)  # formulas kept, cached values dropped on save
    wb.save(stripped)
    profile = (work / "profile").as_uri()
    cmd = [SOFFICE, f"-env:UserInstallation={profile}", "--headless", "--calc",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(work / "out"),
           str(stripped)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=RECALC_TIMEOUT)
    out = work / "out" / "book.xlsx"
    assert res.returncode == 0 and out.exists(), f"LibreOffice failed: {res.stderr[-500:]}"
    return out


# ---- session fixtures -------------------------------------------------------------------
@pytest.fixture(scope="session")
def gt():
    return json.loads(GT_PATH.read_text())


@pytest.fixture(scope="session")
def gtv(gt):
    return {i: v["value"] for i, v in gt["values"].items()}


@pytest.fixture(scope="session")
def raw_assumptions():
    return json.loads((APP / "assumptions.json").read_text())


@pytest.fixture(scope="session")
def market():
    return json.loads((APP / "market_data.json").read_text())


@pytest.fixture(scope="session")
def base_assump(raw_assumptions):
    return rm.assumptions_from_json(raw_assumptions)


@pytest.fixture(scope="session")
def ref_model(gtv, base_assump, market):
    return rm.model_from(gtv, base_assump, market["last_price"])


@pytest.fixture(scope="session")
def ref(ref_model):
    return rm.expected_outputs(ref_model)


@pytest.fixture(scope="session")
def recalced():
    return recalc_book(XLSX)


@pytest.fixture(scope="session")
def vals_wb(recalced):
    return load_workbook(recalced, data_only=True)


@pytest.fixture(scope="session")
def outputs(vals_wb):
    return read_outputs(vals_wb)


@pytest.fixture(scope="session")
def summary():
    return json.loads(SUMMARY.read_text())


def _skip_if_illustrative(gt):
    if gt["_meta"].get("illustrative"):
        pytest.skip("illustrative ground truth: filing-based checks need the real ground truth")


# =========================================================================================
# 0. Ground-truth sanity (guards the verifier itself)
# =========================================================================================
def test_V00_ground_truth_complete_and_real(gt):
    assert set(gt["values"]) == set(rm.required_ids()), "ground truth ids != canonical ids"
    for i, v in gt["values"].items():
        assert isnum(v["value"]), i
    if gt["_meta"].get("illustrative"):
        assert os.environ.get("ALLOW_ILLUSTRATIVE_GT") == "1", \
            "ground_truth.json is illustrative; replace it with the real extracted values"


# =========================================================================================
# 1. Artifacts and integrity
# =========================================================================================
def test_V01_output_files_exist_and_valid():
    assert XLSX.exists() and XLSX.stat().st_size > 0, "nke_dcf.xlsx missing or empty"
    assert SUMMARY.exists() and SUMMARY.stat().st_size > 0, "summary.json missing or empty"
    assert XLSX.stat().st_size < 16 * 1024 * 1024
    assert zipfile.is_zipfile(XLSX), "nke_dcf.xlsx is not a valid xlsx"
    with zipfile.ZipFile(XLSX) as z:
        assert not any("vbaProject" in n for n in z.namelist()), "macros are not allowed"


def test_V02_required_sheets_present():
    wb = load_workbook(XLSX)
    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    assert not missing, f"missing sheets: {missing}"


def test_V03_summary_json_schema(summary):
    for k in rm.SUMMARY_MAP:
        assert k in summary, f"summary.json missing {k}"
        assert isnum(summary[k]), f"summary.json {k} must be a finite number"


def test_V04_protected_files_unmodified(gt):
    hashes = gt.get("protected_hashes", {})
    assert hashes, "ground truth must list protected file hashes"
    for rel, expected in hashes.items():
        p = APP / rel
        assert p.exists(), f"protected file missing: {rel}"
        assert sha256(p) == expected, f"protected file modified: {rel}"


def test_V05_no_external_links_or_network_formulas():
    with zipfile.ZipFile(XLSX) as z:
        assert not any("externalLink" in n for n in z.namelist()), "external links present"
    wb = load_workbook(XLSX)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    low = c.value.lower()
                    assert "http" not in low and "[" not in low, f"{ws.title}!{c.coordinate}"


def test_V06_function_whitelist():
    wb = load_workbook(XLSX)
    used = set()
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    used |= set(re.findall(r"([A-Za-z][A-Za-z0-9_.]*)\s*\(", c.value.upper()))
    assert used <= ALLOWED_FUNCS, f"disallowed functions: {sorted(used - ALLOWED_FUNCS)}"


# =========================================================================================
# 2. Inputs extraction (against hidden ground truth)
# =========================================================================================
@pytest.fixture(scope="session")
def inputs_table(vals_wb):
    ws = vals_wb["Inputs"]
    header = [ws.cell(1, c).value for c in range(1, 5)]
    rows = {}
    for r in range(2, ws.max_row + 1):
        i = ws.cell(r, 1).value
        if i is None:
            continue
        rows.setdefault(i, []).append(
            {"value": ws.cell(r, 2).value, "source": ws.cell(r, 3).value, "page": ws.cell(r, 4).value})
    return header, rows


def test_V10_inputs_schema(inputs_table):
    header, rows = inputs_table
    assert header == ["id", "value", "source_file", "page"], f"bad Inputs header: {header}"
    dups = [i for i, v in rows.items() if len(v) > 1]
    assert not dups, f"duplicate ids: {dups}"
    missing = [i for i in rm.required_ids() if i not in rows]
    assert not missing, f"missing ids: {missing}"


@pytest.mark.parametrize("input_id", rm.required_ids())
def test_V11_input_value_matches_ground_truth(input_id, inputs_table, gtv):
    _, rows = inputs_table
    assert input_id in rows, f"{input_id} missing"
    act = rows[input_id][0]["value"]
    assert isnum(act), f"{input_id}: value is not numeric ({act!r})"
    assert abs(act - gtv[input_id]) <= TOL_INPUT_ABS, f"{input_id}: got {act}, expected {gtv[input_id]}"


@pytest.mark.parametrize("input_id", rm.required_ids())
def test_V12_source_file_is_valid(input_id, inputs_table, gt):
    _skip_if_illustrative(gt)
    src = inputs_table[1][input_id][0]["source"]
    assert src in gt["values"][input_id]["sources"], f"{input_id}: source {src!r} not accepted"
    assert (APP / "filings" / src).exists()


def _printed_forms(v):
    av = abs(v)
    if float(av).is_integer():
        return {f"{int(av):,}", str(int(av))}
    return {f"{av:,.1f}", f"{av:.1f}", f"{av:,.2f}", f"{av:.2f}"}


_PAGE_CACHE = {}


def _page_text(pdf, page):
    key = (str(pdf), page)
    if key not in _PAGE_CACHE:
        import pdfplumber
        with pdfplumber.open(pdf) as doc:
            _PAGE_CACHE[key] = doc.pages[page - 1].extract_text() or ""
    return _PAGE_CACHE[key]


@pytest.mark.parametrize("input_id", rm.required_ids())
def test_V13_cited_page_contains_value(input_id, inputs_table, gt, gtv):
    _skip_if_illustrative(gt)
    row = inputs_table[1][input_id][0]
    page = row["page"]
    assert isinstance(page, (int, float)) and float(page).is_integer() and page >= 1, \
        f"{input_id}: page must be a 1-based integer"
    text = _page_text(APP / "filings" / row["source"], int(page))
    assert any(re.search(r"(?<![\d,.])" + re.escape(f) + r"(?![\d,]|\.\d)", text)
               for f in _printed_forms(gtv[input_id])), \
        f"{input_id}: value {gtv[input_id]} not found on page {int(page)} of {row['source']}"


def test_V14_diag_page_is_primary_statement(inputs_table, gt):
    bad = []
    for i, v in gt["values"].items():
        primary = v.get("primary_pages")
        if primary and inputs_table[1].get(i, [{}])[0].get("page") not in primary:
            bad.append(i)
    diag("V14", not bad, f"{len(bad)} ids cite a non-primary page: {bad[:10]}")


# =========================================================================================
# 3. Assumptions fidelity
# =========================================================================================
def test_V20_assumptions_sheet_matches_json(base_assump, market):
    ws = load_workbook(XLSX, data_only=False)["Assumptions"]
    sheet = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    expected = dict(base_assump, market_price=market["last_price"])
    for k in ASSUMPTION_KEYS:
        assert k in sheet, f"Assumptions sheet missing {k}"
        assert isnum(sheet[k]) and abs(sheet[k] - expected[k]) <= 1e-12, f"{k}: {sheet[k]} != {expected[k]}"


# =========================================================================================
# 4. Recalculation and structure
# =========================================================================================
def test_V30_recalculation_succeeds(recalced):
    assert Path(recalced).exists()


def test_V31_no_error_cells_and_outputs_numeric(vals_wb, outputs):
    errors = []
    for ws in vals_wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERROR_PREFIXES):
                    errors.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not errors, errors[:10]
    for key, ncols in rm.OUTPUT_SPEC.items():
        v = outputs[key]
        assert v is not None, f"label not found (or duplicated): {key}"
        vals = v if ncols else [v]
        assert any(isnum(x) for x in vals), f"{key}: no numeric values"


def test_V32_diag_output_cells_are_formulas():
    wb = load_workbook(XLSX, data_only=False)
    non_formula = []
    for (sheet, label), ncols in rm.OUTPUT_SPEC.items():
        ws = wb[sheet]
        row = find_row(ws, label)
        for i in range(max(ncols, 1)):
            v = ws.cell(row, 2 + i).value
            if v is not None and not (isinstance(v, str) and v.startswith("=")):
                non_formula.append(f"{sheet}!{label}[{i}]")
    diag("V32", not non_formula, non_formula[:10])


def test_V33_diag_cached_values_match_recalc(outputs):
    cached = read_outputs(load_workbook(XLSX, data_only=True))
    stale = [k for k, v in cached.items() if compare({k: v}, {k: outputs[k]}, 1e-9, [k])]
    diag("V33", not stale, f"{len(stale)} stale or missing cached cells: {stale[:10]}")


def test_V34_diag_no_hidden_sheets():
    wb = load_workbook(XLSX)
    hidden = [ws.title for ws in wb if ws.sheet_state != "visible"]
    diag("V34", not hidden, hidden)


# =========================================================================================
# 5-8. Values vs reference model (recalculated workbook, ground-truth inputs)
# =========================================================================================
HIST_KEYS = [k for k in rm.OUTPUT_SPEC if k[0] == "Historicals"]
WACC_KEYS = [k for k in rm.OUTPUT_SPEC if k[0] == "WACC"]
ROW_KEYS = [k for k in rm.OUTPUT_SPEC if k[0] == "DCF" and rm.OUTPUT_SPEC[k] > 0]
BRIDGE_KEYS_ = [k for k in rm.OUTPUT_SPEC if k[0] == "DCF" and rm.OUTPUT_SPEC[k] == 0]
_ids = lambda keys: [f"{s}:{l}" for s, l in keys]


@pytest.mark.parametrize("key", HIST_KEYS, ids=_ids(HIST_KEYS))
def test_V4x_historicals(key, outputs, ref):
    assert compare(outputs, ref, keys=[key]) == []


@pytest.mark.parametrize("key", WACC_KEYS, ids=_ids(WACC_KEYS))
def test_V5x_wacc_components(key, outputs, ref):
    assert compare(outputs, ref, keys=[key]) == []


def test_V54_wacc_within_10bps(outputs, ref):
    w = outputs[("WACC", "WACC")]
    assert isnum(w) and abs(w - ref[("WACC", "WACC")]) <= TOL_WACC_ABS


def test_V53_weights_sum_to_one(outputs):
    assert close(outputs[("WACC", "Equity weight")] + outputs[("WACC", "Debt weight")], 1.0, 1e-9)


def test_V55_wacc_exceeds_terminal_growth(outputs, base_assump):
    assert outputs[("WACC", "WACC")] > base_assump["terminal_growth_rate"]


@pytest.mark.parametrize("key", ROW_KEYS, ids=_ids(ROW_KEYS))
def test_V6x_dcf_rows(key, outputs, ref):
    assert compare(outputs, ref, keys=[key]) == []


def test_V60_growth_path_endpoints(outputs, base_assump):
    g = outputs[("DCF", "Revenue growth")]
    assert close(g[0], outputs[("Historicals", "Revenue CAGR")], 1e-9)
    assert close(g[-1], base_assump["terminal_growth_rate"], 1e-9)


@pytest.mark.parametrize("key", BRIDGE_KEYS_, ids=_ids(BRIDGE_KEYS_))
def test_V7x_valuation_bridge(key, outputs, ref):
    assert compare(outputs, ref, keys=[key]) == []


# =========================================================================================
# 9. Liveness and correctness by perturbation
# =========================================================================================
def perturbation_cases(gtv, price):
    ids = sorted(gtv)
    return {
        "V80_global_inputs": ({i: gtv[i] * (1 + 0.01 * ((k % 7) - 3)) for k, i in enumerate(ids)}, {}),
        "V81_revenue_fy0_x1.10": ({"revenue|FY0": gtv["revenue|FY0"] * 1.10}, {}),
        "V82_ytd_revenue_plus500": ({"revenue|YTD_CUR": gtv["revenue|YTD_CUR"] + 500}, {}),
        "V83_cash_fy0_plus1000": ({"cash|FY0": gtv["cash|FY0"] + 1000}, {}),
        "V84_debt_fy0_plus500": ({"long_term_debt|FY0": gtv["long_term_debt|FY0"] + 500}, {}),
        "V85_q_debt_plus2000": ({"notes_payable|Q_END": gtv["notes_payable|Q_END"] + 2000}, {}),
        "V86_diluted_shares_x1.05": ({"diluted_shares|YTD_CUR": gtv["diluted_shares|YTD_CUR"] * 1.05}, {}),
        "V88_tax_fy-1_x1.1": ({"tax_expense|FY-1": gtv["tax_expense|FY-1"] * 1.1}, {}),
        "V87a_risk_free_rate": ({}, {"risk_free_rate": 0.047}),
        "V87b_beta": ({}, {"beta": 1.3}),
        "V87c_terminal_growth": ({}, {"terminal_growth_rate": 0.03}),
        "V87d_market_price": ({}, {"market_price": price * 1.2}),
    }


def run_perturbation(input_changes, assump_changes):
    wb = load_workbook(XLSX)  # formulas preserved
    ws_in, ws_as = wb["Inputs"], wb["Assumptions"]
    in_rows = {ws_in.cell(r, 1).value: r for r in range(2, ws_in.max_row + 1)}
    as_rows = {ws_as.cell(r, 1).value: r for r in range(2, ws_as.max_row + 1)}
    for i, v in input_changes.items():
        assert i in in_rows, f"Inputs id not found: {i}"
        ws_in.cell(in_rows[i], 2).value = v
    for k, v in assump_changes.items():
        assert k in as_rows, f"Assumptions key not found: {k}"
        ws_as.cell(as_rows[k], 2).value = v
    tmp = Path(tempfile.mkdtemp(prefix="nke_perturb_")) / "perturbed.xlsx"
    wb.save(tmp)
    return read_outputs(load_workbook(recalc_book(tmp), data_only=True))


PERTURBATION_NAMES = list(perturbation_cases({i: 1.0 for i in rm.required_ids()}, 1.0))


@pytest.mark.parametrize("name", PERTURBATION_NAMES)
def test_V8x_perturbation_matches_reference(name, gtv, base_assump, market, outputs):
    input_changes, assump_changes = perturbation_cases(gtv, market["last_price"])[name]
    actual = run_perturbation(input_changes, assump_changes)
    values = dict(gtv, **input_changes)
    a = dict(base_assump, **{k: v for k, v in assump_changes.items() if k != "market_price"})
    price = assump_changes.get("market_price", market["last_price"])
    expected = rm.expected_outputs(rm.model_from(values, a, price))
    bad = compare(actual, expected)
    assert not bad, f"{name}: {len(bad)} outputs do not respond as the reference predicts: {bad[:5]}"

    ev0, eq0 = outputs[("DCF", "Enterprise value")], outputs[("DCF", "Equity value")]
    ev1, eq1 = actual[("DCF", "Enterprise value")], actual[("DCF", "Equity value")]
    if name == "V83_cash_fy0_plus1000":
        assert close(eq1 - eq0, 1000, 1e-9, 1e-6) and close(ev1, ev0, 1e-9)
    if name == "V84_debt_fy0_plus500":
        assert close(eq1 - eq0, -500, 1e-9, 1e-6)
        assert close(actual[("WACC", "WACC")], outputs[("WACC", "WACC")], 1e-12)


# =========================================================================================
# 10. summary.json consistency
# =========================================================================================
@pytest.mark.parametrize("field", list(rm.SUMMARY_MAP))
def test_V90_summary_matches_recalculated_workbook(field, summary, outputs):
    cell = outputs[rm.SUMMARY_MAP[field]]
    assert isnum(cell) and isnum(summary[field])
    assert close(summary[field], cell, TOL_JSON_REL), f"{field}: json {summary[field]} vs workbook {cell}"


def test_V91_upside_identity(summary):
    assert close(summary["upside_pct"], summary["value_per_share"] / summary["market_price"] - 1, 1e-9)


def test_V92_summary_matches_reference(summary, ref_model):
    exp = rm.summary_from(ref_model)
    assert abs(summary["wacc"] - exp["wacc"]) <= TOL_WACC_ABS
    for k in ("enterprise_value", "equity_value", "value_per_share"):
        assert close(summary[k], exp[k], TOL_E2E_REL), f"{k}: {summary[k]} vs {exp[k]}"


def test_V93_market_price_from_fixed_file(summary, market):
    assert summary["market_price"] == market["last_price"]


# =========================================================================================
# 11. Invariants and sanity
# =========================================================================================
def test_V100_discount_factors(outputs):
    df = outputs[("DCF", "Discount factor")]
    assert all(isnum(x) and 0 < x < 1 for x in df)
    assert all(df[i] > df[i + 1] for i in range(len(df) - 1))


def test_V101_wacc_in_plausible_range(outputs):
    assert 0.04 <= outputs[("WACC", "WACC")] <= 0.15


def test_V102_terminal_value_share_of_ev(outputs):
    share = outputs[("DCF", "PV of terminal value")] / outputs[("DCF", "Enterprise value")]
    assert 0.5 <= share <= 0.9, f"PV(TV)/EV = {share:.1%}"


def test_V103_internal_identities(outputs):
    o = lambda s, l: outputs[(s, l)]
    assert close(o("DCF", "Enterprise value"), o("DCF", "Sum of PV of FCF") + o("DCF", "PV of terminal value"))
    assert close(sum(o("DCF", "PV of FCF")), o("DCF", "Sum of PV of FCF"))
    assert close(o("DCF", "Equity value"), o("DCF", "Enterprise value") + o("DCF", "Cash") - o("DCF", "Debt"))
    assert close(o("DCF", "Value per share") * o("DCF", "Diluted shares"), o("DCF", "Equity value"))
    assert close(o("DCF", "PV of terminal value"),
                 o("DCF", "Terminal value") * o("DCF", "Discount factor")[-1])


def test_V104_diag_value_vs_market(summary):
    diag("V104", True, f"value/share {summary['value_per_share']:.2f} vs market {summary['market_price']:.2f} "
                       f"(upside {summary['upside_pct']:.1%}); informational only")
