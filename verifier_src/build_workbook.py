"""Golden-solution workbook builder for nke-dcf-valuation.

Takes the extracted inputs (the PDF-extraction half of the solution produces this JSON),
the pinned assumptions and the fixed market price, then:
  1. writes nke_dcf.xlsx with live formulas (Assumptions, Inputs, WACC, Historicals, DCF),
  2. recalculates it with LibreOffice headless and keeps the recalculated file,
  3. writes summary.json from the recalculated cells.

Usage:
  python build_workbook.py --inputs inputs.json --assumptions /app/assumptions.json \
      --market /app/market_data.json --out-dir /app/output

inputs.json is ``{"values": {id: {"value": x, "source_file": "...", "page": n}}}``
(``sources`` / ``pages`` lists, as in ground_truth.json, are also accepted).
"""
import argparse
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FY = ["FY-2", "FY-1", "FY0"]
FY_KEYS = ["revenue", "pretax_income", "interest_net", "tax_expense", "depreciation", "capex",
           "receivables", "inventories", "prepaid", "payables", "accrued"]
YTD_KEYS = ["revenue", "pretax_income", "interest_net", "tax_expense"]
BRIDGE_KEYS = ["cash", "short_term_investments", "notes_payable", "current_ltd", "long_term_debt"]
DEBT_KEYS = ["notes_payable", "current_ltd", "long_term_debt"]


def input_ids():
    ids = [f"{k}|{p}" for p in FY for k in FY_KEYS]
    ids += [f"{k}|FY0" for k in BRIDGE_KEYS]
    ids += [f"{k}|{p}" for p in ("YTD_CUR", "YTD_PRI") for k in YTD_KEYS]
    ids += ["diluted_shares|YTD_CUR"]
    ids += [f"{k}|Q_END" for k in DEBT_KEYS]
    return ids


def I(id_):
    """Formula fragment that looks an input up by id."""
    return f'INDEX(Inputs!$B:$B,MATCH("{id_}",Inputs!$A:$A,0))'


def build(values, raw_assumptions, price, path):
    wb = Workbook()
    bold = Font(bold=True)

    # ---- Assumptions ------------------------------------------------------------------
    ws = wb.active
    ws.title = "Assumptions"
    ws.append(["key", "value"])
    arows = [
        ("risk_free_rate", raw_assumptions["wacc"]["risk_free_rate"]),
        ("equity_risk_premium", raw_assumptions["wacc"]["equity_risk_premium"]),
        ("beta", raw_assumptions["wacc"]["beta"]),
        ("pre_tax_cost_of_debt", raw_assumptions["wacc"]["pre_tax_cost_of_debt"]),
        ("terminal_growth_rate", raw_assumptions["terminal_value"]["growth_rate"]),
        ("explicit_forecast_years", raw_assumptions["forecast"]["explicit_forecast_years"]),
        ("market_price", price),
    ]
    A = {}
    for r, (k, v) in enumerate(arows, start=2):
        ws.cell(r, 1, k)
        ws.cell(r, 2, v)
        A[k] = f"Assumptions!$B${r}"

    # ---- Inputs -----------------------------------------------------------------------
    wi = wb.create_sheet("Inputs")
    wi.append(["id", "value", "source_file", "page"])
    for id_ in input_ids():
        e = values[id_]
        src = e.get("source_file") or (e.get("sources") or [None])[0]
        page = e.get("page") or (e.get("pages") or [None])[0]
        wi.append([id_, e["value"], src, page])

    # ---- Historicals ------------------------------------------------------------------
    wh = wb.create_sheet("Historicals")
    wh.append(["Label"] + FY)
    labels = {2: "Revenue", 3: "EBIT", 4: "Revenue growth", 5: "EBIT margin",
              6: "Effective tax rate", 7: "D&A % of revenue", 8: "Capex % of revenue",
              9: "NWC", 10: "NWC % of revenue", 11: "Net investment"}
    for r, l in labels.items():
        wh.cell(r, 1, l)
    cols = ["B", "C", "D"]
    for j, (c, p) in enumerate(zip(cols, FY)):
        wh[f"{c}2"] = f"={I('revenue|' + p)}"
        wh[f"{c}3"] = f"={I('pretax_income|' + p)}+{I('interest_net|' + p)}"
        if j > 0:
            wh[f"{c}4"] = f"={c}2/{cols[j - 1]}2-1"
        wh[f"{c}5"] = f"={c}3/{c}2"
        wh[f"{c}6"] = f"={I('tax_expense|' + p)}/{I('pretax_income|' + p)}"
        wh[f"{c}7"] = f"={I('depreciation|' + p)}/{c}2"
        wh[f"{c}8"] = f"={I('capex|' + p)}/{c}2"
        wh[f"{c}9"] = (f"={I('receivables|' + p)}+{I('inventories|' + p)}+{I('prepaid|' + p)}"
                       f"-{I('payables|' + p)}-{I('accrued|' + p)}")
        wh[f"{c}10"] = f"={c}9/{c}2"
        if j > 0:
            wh[f"{c}11"] = (f"={I('capex|' + p)}-{I('depreciation|' + p)}+({c}9-{cols[j - 1]}9)")
    scalars = [(13, "Avg EBIT margin", "=AVERAGE(B5:D5)"),
               (14, "Avg effective tax rate", "=AVERAGE(B6:D6)"),
               (15, "Avg D&A % of revenue", "=AVERAGE(B7:D7)"),
               (16, "Avg capex % of revenue", "=AVERAGE(B8:D8)"),
               (17, "Avg NWC % of revenue", "=AVERAGE(B10:D10)"),
               (18, "Revenue CAGR", "=POWER(D2/B2,1/2)-1"),
               (19, "LTM revenue", f"=D2+{I('revenue|YTD_CUR')}-{I('revenue|YTD_PRI')}")]
    for r, l, f in scalars:
        wh.cell(r, 1, l)
        wh.cell(r, 2, f)

    # ---- WACC -------------------------------------------------------------------------
    ww = wb.create_sheet("WACC")
    ww.append(["Label", "Value"])
    debt_q = "+".join(I(f"{k}|Q_END") for k in DEBT_KEYS)
    wrows = [
        (2, "Cost of equity", f"={A['risk_free_rate']}+{A['beta']}*{A['equity_risk_premium']}"),
        (3, "After-tax cost of debt", f"={A['pre_tax_cost_of_debt']}*(1-Historicals!$B$14)"),
        (4, "Market cap", f"={A['market_price']}*{I('diluted_shares|YTD_CUR')}"),
        (5, "Debt", f"={debt_q}"),
        (6, "Equity weight", "=B4/(B4+B5)"),
        (7, "Debt weight", "=B5/(B4+B5)"),
        (8, "WACC", "=B6*B2+B7*B3"),
    ]
    for r, l, f in wrows:
        ww.cell(r, 1, l)
        ww.cell(r, 2, f)

    # ---- DCF --------------------------------------------------------------------------
    wd = wb.create_sheet("DCF")
    wd.cell(1, 1, "Label")
    for t in range(1, 11):
        wd.cell(1, 1 + t, t)
    series = {2: "Revenue growth", 3: "Revenue", 4: "EBIT", 5: "NOPAT", 6: "Capex", 7: "D&A",
              8: "NWC", 9: "Change in NWC", 10: "Net investment", 11: "Free cash flow",
              12: "Discount factor", 13: "PV of FCF"}
    for r, l in series.items():
        wd.cell(r, 1, l)
    gT, yrs = A["terminal_growth_rate"], A["explicit_forecast_years"]
    for t in range(1, 11):
        L = get_column_letter(1 + t)
        P = get_column_letter(t)
        wd[f"{L}2"] = f"=Historicals!$B$18+({gT}-Historicals!$B$18)*({L}$1-1)/({yrs}-1)"
        wd[f"{L}3"] = (f"=Historicals!$B$19*(1+{L}2)" if t == 1 else f"={P}3*(1+{L}2)")
        wd[f"{L}4"] = f"=Historicals!$B$13*{L}3"
        wd[f"{L}5"] = f"={L}4*(1-Historicals!$B$14)"
        wd[f"{L}6"] = f"=Historicals!$B$16*{L}3"
        wd[f"{L}7"] = f"=Historicals!$B$15*{L}3"
        wd[f"{L}8"] = f"=Historicals!$B$17*{L}3"
        wd[f"{L}9"] = (f"={L}8-Historicals!$B$17*Historicals!$B$19" if t == 1 else f"={L}8-{P}8")
        wd[f"{L}10"] = f"={L}6-{L}7+{L}9"
        wd[f"{L}11"] = f"={L}5-{L}10"
        wd[f"{L}12"] = f"=1/POWER(1+WACC!$B$8,{L}$1)"
        wd[f"{L}13"] = f"={L}11*{L}12"
    debt_k = "+".join(I(f"{k}|FY0") for k in DEBT_KEYS)
    drows = [
        (15, "Sum of PV of FCF", "=SUM(B13:K13)"),
        (16, "Terminal value", f"=K11*(1+{gT})/(WACC!$B$8-{gT})"),
        (17, "PV of terminal value", "=B16*K12"),
        (18, "Enterprise value", "=B15+B17"),
        (19, "Cash", f"={I('cash|FY0')}+{I('short_term_investments|FY0')}"),
        (20, "Debt", f"={debt_k}"),
        (21, "Equity value", "=B18+B19-B20"),
        (22, "Diluted shares", f"={I('diluted_shares|YTD_CUR')}"),
        (23, "Value per share", "=B21/B22"),
        (24, "Market price", f"={A['market_price']}"),
        (25, "Upside", "=B23/B24-1"),
    ]
    for r, l, f in drows:
        wd.cell(r, 1, l)
        wd.cell(r, 2, f)

    for sheet in wb:
        for c in sheet[1]:
            c.font = bold
        sheet.column_dimensions["A"].width = 28
    wb.save(path)


def recalc(src, dst):
    """Recalculate with LibreOffice headless and write the recalculated copy to dst."""
    work = Path(tempfile.mkdtemp(prefix="nke_build_"))
    soffice = os.environ.get("SOFFICE", "soffice")
    cmd = [soffice, f"-env:UserInstallation={(work / 'profile').as_uri()}", "--headless",
           "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(work), str(src)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    out = work / Path(src).name
    if res.returncode != 0 or not out.exists():
        raise RuntimeError(f"LibreOffice recalculation failed: {res.stderr[-500:]}")
    shutil.copy(out, dst)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inputs", required=True)
    ap.add_argument("--assumptions", required=True)
    ap.add_argument("--market", required=True)
    ap.add_argument("--out-dir", required=True)
    args = ap.parse_args()

    raw = json.loads(Path(args.inputs).read_text())
    values = raw.get("values", raw)
    assumptions = json.loads(Path(args.assumptions).read_text())
    price = json.loads(Path(args.market).read_text())["last_price"]
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    draft = Path(tempfile.mkdtemp(prefix="nke_draft_")) / "nke_dcf.xlsx"
    build(values, assumptions, price, draft)
    final = out_dir / "nke_dcf.xlsx"
    recalc(draft, final)

    wb = load_workbook(final, data_only=True)
    find = lambda ws, label: next(ws.cell(r, 2).value for r in range(1, ws.max_row + 1)
                                  if ws.cell(r, 1).value == label)
    summary = {
        "wacc": find(wb["WACC"], "WACC"),
        "enterprise_value": find(wb["DCF"], "Enterprise value"),
        "equity_value": find(wb["DCF"], "Equity value"),
        "value_per_share": find(wb["DCF"], "Value per share"),
        "market_price": find(wb["DCF"], "Market price"),
        "upside_pct": find(wb["DCF"], "Upside"),
    }
    (out_dir / "summary.json").write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
