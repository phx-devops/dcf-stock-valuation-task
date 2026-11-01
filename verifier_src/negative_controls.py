"""Authoring-time self-test for the verifier (not part of the graded run).

Takes a known-good output directory (golden solution output), applies one targeted defect at a
time, runs test_outputs.py on the broken copy, and checks that the intended check fails.
Also runs the verifier twice on the golden output to confirm it is deterministic.

Usage:
  APP_DIR=/path/to/app ALLOW_ILLUSTRATIVE_GT=1 python negative_controls.py
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
APP = Path(os.environ.get("APP_DIR", "/app"))
GOLDEN = Path(os.environ.get("GOLDEN_DIR", str(APP / "output")))


def run_verifier(out_dir):
    env = dict(os.environ, OUT_DIR=str(out_dir), APP_DIR=str(APP),
               REPORT_PATH=str(Path(out_dir) / "report.json"))
    res = subprocess.run([sys.executable, "-m", "pytest", str(HERE / "test_outputs.py"), "-q",
                          "-p", "no:cacheprovider", "--tb=no", "-rf"],
                         capture_output=True, text=True, env=env, timeout=900)
    failed = sorted(set(re.findall(r"^FAILED \S+?::(.+?)(?: - .*)?$", res.stdout, flags=re.M)))
    summary = re.findall(r"(\d+ (?:passed|failed).*?)(?: in [\d.]+s.*)?$", res.stdout, flags=re.M)
    return res.returncode, failed, summary[-1] if summary else res.stdout[-300:]


def edit_xlsx(out_dir, fn):
    path = Path(out_dir) / "nke_dcf.xlsx"
    wb = load_workbook(path)
    fn(wb)
    wb.save(path)


def row_of(ws, label):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == label)


def hardcode(out_dir):
    cached = load_workbook(Path(out_dir) / "nke_dcf.xlsx", data_only=True)

    def fn(wb):
        for name in ("WACC", "Historicals", "DCF"):
            for row in wb[name].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        c.value = cached[name][c.coordinate].value
    edit_xlsx(out_dir, fn)


def tv_year9(out_dir):
    def fn(wb):
        ws = wb["DCF"]
        r_tv, r_pv = row_of(ws, "Terminal value"), row_of(ws, "PV of terminal value")
        ws.cell(r_tv, 2).value = ws.cell(r_tv, 2).value.replace("K11", "J11")
        ws.cell(r_pv, 2).value = ws.cell(r_pv, 2).value.replace("K12", "J12")
    edit_xlsx(out_dir, fn)


def wrong_ltm(out_dir):
    def fn(wb):
        ws = wb["Historicals"]
        ws.cell(row_of(ws, "LTM revenue"), 2).value = "=D2"
    edit_xlsx(out_dir, fn)


def ebit_operating(out_dir):
    def fn(wb):
        ws = wb["Historicals"]
        r = row_of(ws, "EBIT")
        for c in (2, 3, 4):
            ws.cell(r, c).value = ws.cell(r, c).value.split("+INDEX")[0]
    edit_xlsx(out_dir, fn)


def _set_input(wb, input_id, delta):
    ws = wb["Inputs"]
    r = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == input_id)
    ws.cell(r, 2).value = ws.cell(r, 2).value + delta


def basic_shares(out_dir):
    edit_xlsx(out_dir, lambda wb: _set_input(wb, "diluted_shares|YTD_CUR", -20))


def leases_in_debt(out_dir):
    edit_xlsx(out_dir, lambda wb: _set_input(wb, "long_term_debt|FY0", 2000))


def assumption_tampered(out_dir):
    def fn(wb):
        ws = wb["Assumptions"]
        ws.cell(row_of(ws, "beta"), 2).value = 1.0
    edit_xlsx(out_dir, fn)


def pre_restatement(out_dir):
    """accrued|FY-1 reverts to the pre-restatement figure (5,916 -> 5,911, a -5 shift)."""
    edit_xlsx(out_dir, lambda wb: _set_input(wb, "accrued|FY-1", -5))


def fake_pages(out_dir):
    """Correct values but every citation points at page 1 (memorized figures, no real source)."""
    def fn(wb):
        ws = wb["Inputs"]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 4).value = 1
    edit_xlsx(out_dir, fn)


REAL_ONLY = {"fake_pages"}  # needs real (non-illustrative) ground truth and the filing PDFs


def _edit_summary(out_dir, key, factor):
    p = Path(out_dir) / "summary.json"
    d = json.loads(p.read_text())
    d[key] = d[key] * factor
    p.write_text(json.dumps(d))


CONTROLS = {
    "hardcoded_values": (hardcode, ["test_V8x_perturbation_matches_reference[V80_global_inputs]"]),
    "terminal_value_year9": (tv_year9, ["test_V7x_valuation_bridge[DCF:Terminal value]"]),
    "wrong_ltm_base": (wrong_ltm, ["test_V4x_historicals[Historicals:LTM revenue]"]),
    "ebit_is_pretax_only": (ebit_operating, ["test_V4x_historicals[Historicals:EBIT]"]),
    "basic_not_diluted_shares": (basic_shares, ["test_V11_input_value_matches_ground_truth[diluted_shares|YTD_CUR]"]),
    "leases_in_debt": (leases_in_debt, ["test_V11_input_value_matches_ground_truth[long_term_debt|FY0]"]),
    "assumption_tampered": (assumption_tampered, ["test_V20_assumptions_sheet_matches_json"]),
    "pre_restatement_value": (pre_restatement, ["test_V11_input_value_matches_ground_truth[accrued|FY-1]"]),
    "fake_pages": (fake_pages, ["test_V13_cited_page_contains_value[revenue|FY0]"]),
    "summary_mismatch": (lambda d: _edit_summary(d, "value_per_share", 1.1),
                         ["test_V90_summary_matches_recalculated_workbook[value_per_share]"]),
    "market_price_not_fixed": (lambda d: _edit_summary(d, "market_price", 1.01),
                               ["test_V93_market_price_from_fixed_file"]),
}


def main():
    ok = True
    rc1, f1, s1 = run_verifier(GOLDEN)
    rc2, f2, s2 = run_verifier(GOLDEN)
    det = (rc1, f1, s1) == (rc2, f2, s2)
    good = rc1 == 0 and not f1
    print(f"golden run 1: rc={rc1} failed={f1} [{s1}]")
    print(f"golden passes: {good}; two runs identical: {det}")
    ok &= good and det
    gt_path = Path(os.environ.get("GT_PATH", str(HERE / "ground_truth.json")))
    illustrative = json.loads(gt_path.read_text())["_meta"].get("illustrative", False)
    for name, (mutate, expected) in CONTROLS.items():
        if name in REAL_ONLY and illustrative:
            print(f"[SKIP] {name}: needs the real ground truth and filing PDFs (page checks are skipped on illustrative data)")
            continue
        work =Path(tempfile.mkdtemp(prefix=f"nc_{name}_"))
        shutil.copy(GOLDEN / "nke_dcf.xlsx", work / "nke_dcf.xlsx")
        shutil.copy(GOLDEN / "summary.json", work / "summary.json")
        mutate(work)
        rc, failed, summary = run_verifier(work)
        hit = all(e in failed for e in expected)
        print(f"[{'OK ' if rc != 0 and hit else 'BAD'}] {name}: rc={rc}, {len(failed)} failing "
              f"checks, intended check failed={hit}")
        if not hit:
            print("      expected:", expected, "\n      got:", failed[:8])
        ok &= rc != 0 and hit
    print("ALL NEGATIVE CONTROLS BEHAVE AS EXPECTED" if ok else "SELF-TEST FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())